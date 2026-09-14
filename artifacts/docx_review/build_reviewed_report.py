from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path("artifacts/CustomTools-VI_v1_bao_cao_sau_review.docx")

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "DCE6F1"
PALE_GREEN = "E2F0D9"
LIGHT = "F3F6F9"
WHITE = "FFFFFF"
TEXT = "1F2937"
MUTED = "5B6573"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color: str = "B7C5D5", size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def set_repeat_font(run, name: str = "Aptos") -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Cm(widths_cm[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)
    for ridx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(values):
            cell = cells[idx]
            cell.width = Cm(widths_cm[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ridx % 2:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc: Document, title: str, body: str, fill: str = PALE_GREEN) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.25)
    p.paragraph_format.right_indent = Cm(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), BLUE)
    left.set(qn("w:space"), "8")
    borders.append(left)
    p_pr.append(borders)
    run = p.add_run(f"{title}: ")
    run.bold = True
    p.add_run(body)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, color in (("Title", 24, NAVY), ("Subtitle", 11, MUTED), ("Heading 1", 16, NAVY), ("Heading 2", 12.5, BLUE), ("Heading 3", 11, TEXT)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Subtitle" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(4)
    title.add_run("BÁO CÁO THỐNG KÊ VÀ KIỂM ĐỊNH\nCUSTOMTOOLS-VI V1")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Bản đã review • Đối chiếu dữ liệu phát hành ngày 12/08/2026")
    line = doc.add_paragraph()
    line.paragraph_format.space_before = Pt(3)
    line.paragraph_format.space_after = Pt(12)
    run = line.add_run("━" * 54)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run.font.size = Pt(8)

    add_callout(doc, "Kết luận", "Bộ dữ liệu gồm 8.000 mẫu và 40 công cụ, đã vượt toàn bộ kiểm tra tự động ở cấp bản ghi và kiểm toán độc lập. Các kết quả này chưa thay thế đánh giá thủ công của con người hoặc thử nghiệm với API thực tế.")

    add_heading(doc, "1. Tổng quan", 1)
    p = doc.add_paragraph()
    p.add_run("CustomTools-VI v1").bold = True
    p.add_run(" là bộ dữ liệu function calling tiếng Việt tổng hợp, tập trung vào ngữ cảnh sử dụng tại Việt Nam. Bộ dữ liệu được tạo theo hướng frame-first, sau đó kiểm tra bằng validator trong bộ sinh và một auditor độc lập trên các tệp JSONL đã tuần tự hóa.")
    add_table(doc, ["Chỉ số", "Kết quả"], [
        ["Tổng số mẫu", "8.000"],
        ["Tổng số công cụ", "40"],
        ["Nhóm chức năng", "10"],
        ["Candidate tools mỗi mẫu", "10"],
        ["Positive", "4.800 (60%)"],
        ["Negative/no-call", "3.200 (40%)"],
        ["Tổng function calls", "5.520"],
        ["Single-call", "4.080 (51% toàn bộ dữ liệu)"],
        ["Parallel two-call", "720 (9% toàn bộ dữ liệu)"],
        ["Multi-call trong positive", "15%"],
    ], [10.2, 6.0])

    add_heading(doc, "2. Phân chia dữ liệu", 1)
    add_table(doc, ["Split", "Positive", "Negative", "Tổng", "Tỷ lệ"], [
        ["train", "3.600", "2.000", "5.600", "70%"],
        ["val_seen", "200", "200", "400", "5%"],
        ["val_unseen", "200", "200", "400", "5%"],
        ["test_seen", "400", "400", "800", "10%"],
        ["test_unseen", "400", "400", "800", "10%"],
        ["Tổng", "4.800", "3.200", "8.000", "100%"],
    ], [4.2, 3.0, 3.0, 3.0, 3.0])
    add_bullet(doc, "Train: 5.600 mẫu (70%).")
    add_bullet(doc, "Validation: 800 mẫu (10%).")
    add_bullet(doc, "Test: 1.600 mẫu (20%).")
    add_bullet(doc, "Seen evaluation: 1.200 mẫu; unseen evaluation: 1.200 mẫu.")

    add_heading(doc, "3. Phân bố nội dung", 1)
    add_heading(doc, "3.1. Nhóm chức năng", 2)
    domains = ["Ẩm thực & Đặc sản", "Du lịch & Địa danh", "Tài chính & Ngân hàng", "Giáo dục Việt Nam", "Y tế & Sức khỏe", "Thương mại điện tử", "Giao thông & Di chuyển", "Hành chính công", "Giải trí & Lịch", "Nông nghiệp & Thời tiết"]
    add_table(doc, ["Nhóm chức năng", "Số mẫu", "Tỷ lệ"], [[x, "800", "10%"] for x in domains], [10.2, 3.0, 3.0])
    add_heading(doc, "3.2. Phong cách câu hỏi", 2)
    add_table(doc, ["Phong cách", "Số mẫu", "Tỷ lệ"], [
        ["Trung tính", "2.800", "35%"], ["Lịch sự", "1.600", "20%"], ["Hội thoại tự nhiên", "1.600", "20%"], ["Nhiễu/viết tắt", "1.200", "15%"], ["Code-mixed", "800", "10%"],
    ], [10.2, 3.0, 3.0])
    doc.add_paragraph("Có 3.686 mẫu (46,08%) được bổ sung ngữ cảnh đa dạng tự nhiên để giảm nguy cơ gần trùng mà không thay đổi gold call.")
    add_heading(doc, "3.3. Negative samples", 2)
    add_table(doc, ["Loại negative", "Số mẫu", "Trong negative", "Toàn bộ"], [
        ["Hard near-miss", "1.280", "40%", "16%"], ["Non-action", "800", "25%", "10%"], ["Out-of-scope", "640", "20%", "8%"], ["Negated/hypothetical", "480", "15%", "6%"], ["Tổng", "3.200", "100%", "40%"],
    ], [7.2, 2.8, 3.2, 3.0])
    doc.add_paragraph("Mỗi hard near-miss có tối thiểu 5 hard distractors thuộc nhóm chức năng chính hoặc nhóm liên quan.")

    add_heading(doc, "4. Tool split và độ phủ positive call", 1)
    add_table(doc, ["Tool tier", "Số tool", "Positive calls mỗi tool"], [["Seen", "20", "241-242"], ["Dev-unseen", "10", "23"], ["Test-unseen", "10", "46"]], [7.2, 3.5, 5.5])
    doc.add_paragraph("Độ phủ được cân bằng trong từng tier, với chênh lệch tối đa một call. Các công cụ unseen có ít hơn 100 positive vì chỉ xuất hiện trong split đánh giá tương ứng; đây là chủ ý nhằm ngăn schema unseen lọt vào tập train. Vì vậy, mục tiêu “100+ positive cho cả 40 tool” trong kế hoạch ban đầu không đồng thời tương thích với thiết kế zero-shot unseen hiện tại.")

    add_heading(doc, "5. Phân bố kiểu tham số", 1)
    doc.add_paragraph("Bốn mươi schema có tổng cộng 144 định nghĩa tham số.")
    add_table(doc, ["Kiểu tham số", "Số lượng", "Tỷ lệ"], [
        ["String không enum", "64", "44,44%"], ["Integer", "31", "21,53%"], ["Number", "6", "4,17%"], ["Integer + Number", "37", "25,69%"], ["Enum", "28", "19,44%"], ["Boolean", "15", "10,42%"],
    ], [9.2, 3.5, 3.5])
    doc.add_paragraph("Tỷ lệ optional parameter vắng mặt trong positive calls là 46,63%, nằm trong mục tiêu 35-50%.")

    add_heading(doc, "6. Kết quả kiểm định sau review", 1)
    add_table(doc, ["Kiểm tra", "Kết quả"], [
        ["Bản ghi vượt validator trong generator", "8.000/8.000"],
        ["Bản ghi vượt auditor độc lập", "8.000/8.000"],
        ["Argument mentions được kiểm tra", "16.364"],
        ["Boolean mentions được kiểm tra ngữ nghĩa", "1.094"],
        ["Surface lặp nhưng vẫn có offset chính xác", "115"],
        ["Argument reconstruction thành công", "8.000/8.000"],
        ["ID trùng", "0"],
        ["Query trùng sau chuẩn hóa", "0"],
        ["Cặp lexical Jaccard ≥ 0,90", "0"],
        ["Cặp lexical gần trùng xuyên split", "0"],
        ["Scenario-family leakage", "0"],
        ["Lỗi tiếng Việt theo pattern đã biết", "0"],
        ["Lỗi audit cuối", "0"],
        ["Manifest integrity", "Passed"],
        ["Test suite", "95/95 passed"],
        ["Ruff", "Passed"],
        ["Mypy", "Passed"],
        ["Tái lập byte-level cùng seed", "Passed"],
    ], [11.5, 4.7])
    p = doc.add_paragraph()
    p.add_run("SHA-256 của generation_manifest.json: ").bold = True
    r = p.add_run("8E9C2840C3278660D271E944DDA7DE8F7DB2FA14D8D7FC60B19B92C4E91CF3C8")
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)

    add_heading(doc, "7. Giới hạn và cách diễn giải", 1)
    add_bullet(doc, "Chưa có human validation hoặc inter-annotator agreement.")
    add_bullet(doc, "Chưa thực thi API thật để kiểm tra kết quả ngoài đời.")
    add_bullet(doc, "Chưa chạy semantic embedding dedup bằng BGE-M3 vì môi trường phát hành không có model weights.")
    add_bullet(doc, "Kết quả 0 near-duplicate chỉ áp dụng cho exact/normalized dedup và exhaustive word-set Jaccard tại ngưỡng 0,90; đây không phải kết quả cosine BGE-M3.")
    add_bullet(doc, "Phiên bản v1 sử dụng parameter schema phẳng, chưa bao gồm array hoặc nested object.")
    add_callout(doc, "Diễn giải đúng", "Bộ dữ liệu là AI-synthesized và được kiểm định xác định bằng mã. Trạng thái “passed” phản ánh các kiểm tra tự động đã công bố, không phải chứng nhận chất lượng ngữ nghĩa bởi chuyên gia con người.", PALE_BLUE)

    add_heading(doc, "8. Tệp phát hành và nguồn đối chiếu", 1)
    files = [
        ("train.jsonl", "5.600 mẫu"), ("val_seen.jsonl", "400 mẫu"), ("val_unseen.jsonl", "400 mẫu"), ("test_seen.jsonl", "800 mẫu"), ("test_unseen.jsonl", "800 mẫu"), ("tools.json", "40 công cụ"), ("qa_report.json", "Báo cáo QA"), ("independent_audit.json", "Kiểm toán độc lập"), ("dataset_card.md", "Mô tả dữ liệu"), ("generation_manifest.json", "Manifest và checksum"),
    ]
    add_table(doc, ["Tệp", "Vai trò/quy mô"], [[a, b] for a, b in files], [9.0, 7.2])
    p = doc.add_paragraph()
    p.add_run("Thư mục phát hành: ").bold = True
    r = p.add_run("data/custom_vi/v1/")
    r.font.name = "Consolas"

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                set_repeat_font(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name is None:
                            set_repeat_font(run)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("CustomTools-VI v1 • Báo cáo đã review • 15/08/2026")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(MUTED)

    core = doc.core_properties
    core.title = "Báo cáo thống kê và kiểm định CustomTools-VI v1"
    core.subject = "Bản đã review và đối chiếu với QA report, independent audit và generation manifest"
    core.author = "Thinh"
    core.keywords = "CustomTools-VI, tool calling, dataset, QA, audit"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
